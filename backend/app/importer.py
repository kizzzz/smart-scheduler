"""排班表导入：CSV / TSV / Excel 走纯确定性解析，图片走视觉模型识别。

三条来源、一个出口：不管数据是从 CSV 读来的还是视觉模型认出来的，最终都收敛成
`List[Slot]`，再由调用方交给同一个 `validator.validate()` 体检。导入功能的核心价值就是
这句体检结论——「把现有排班表拖进来，一键看出有没有违规」，所以绕过校验器的导入是没有
意义的，边界必须守住：

- CSV / Excel **完全不过 LLM**。表格是结构化数据，用模型解析只会引入不确定性；
- 图片来源的识别结果同样不被信任：工号归一、非法 token 判定与合规判定都在确定性代码里，
  模型只负责「看图说出格子里写了什么」；
- 任何无法归一为配置内工号的 token 一律进 unresolved，**绝不猜测映射**。默认配置的员工
  档案没有姓名（name 就是工号），把「小王」映射成某个工号就是凭空编造数据。

配置化后，「哪些日期、哪些班次、哪些工号算合法」不再由 data.py 的常量决定，而是由
`_Dialect` 从当前配置编译出来的别名表决定；解析算法本身一行没动——它面对的是形状问题
（矩阵表还是长表、分隔符是什么），与门店排几天几班无关。
"""
from __future__ import annotations

import base64
import csv
import io
import re
import weakref
from typing import Dict, List, Optional, Tuple

from . import llm
from .config import ConfigIndex, ConfigLike, default_index, index_of
from .models import ImportExtraction, Slot, UnresolvedToken

MAX_UPLOAD_BYTES = 5 * 1024 * 1024

IMAGE_EXTS = {".png", ".jpg", ".jpeg", ".webp"}
TABLE_EXTS = {".csv", ".tsv", ".xlsx", ".xls"}
ALLOWED_EXTS = IMAGE_EXTS | TABLE_EXTS

_MIME_BY_EXT = {".png": "image/png", ".jpg": "image/jpeg", ".jpeg": "image/jpeg", ".webp": "image/webp"}


# ---------- 归一化基础设施 ----------

# 全角 → 半角：Excel 里手打的「Ｅ０１」「，」非常常见，不折平后面所有别名匹配都会漏
_FULLWIDTH = {c: c - 0xFEE0 for c in range(0xFF01, 0xFF5F)}
_FULLWIDTH[0x3000] = 0x20


def _half(s: str) -> str:
    return str(s).translate(_FULLWIDTH)


def _norm(s: object) -> str:
    """列名/值的统一归一：全角折半、去空白、小写。别名表全部按这个形态登记。"""
    return re.sub(r"\s+", "", _half("" if s is None else str(s))).lower()


_DAY_COLUMN_ALIASES = {"日期", "星期", "weekday", "day", "date"}
_SHIFT_COLUMN_ALIASES = {"班次", "班", "shift"}
_EMPLOYEE_COLUMN_ALIASES = {
    "员工", "人员", "工号", "员工工号", "员工编号", "staff", "employees", "employee",
}

_EN_DAYS = [
    ("mon", "monday"),
    ("tue", "tuesday", "tues"),
    ("wed", "wednesday"),
    ("thu", "thursday", "thur", "thurs"),
    ("fri", "friday"),
    ("sat", "saturday"),
    ("sun", "sunday"),
]

# 班次口语别名：只对「配置里真的有这个班次」时才登记，否则默认两班制下「中班」必须仍然
# 是无法识别的班次（识别失败要报给店长，不能悄悄归到晚班）
_SHIFT_SYNONYMS: Dict[str, Tuple[str, ...]] = {
    "早班": ("早", "上午", "上午班", "白班", "早上", "morning", "am"),
    "晚班": ("晚", "下午", "下午班", "夜班", "晚上", "evening", "pm"),
    "中班": ("中", "中午", "午班", "afternoon"),
    "夜班": ("夜", "夜间", "通宵", "night"),
}

# 单元格里的「这格没人」写法。留着不处理会全部变成 unresolved，把真正的问题淹掉
_EMPTY_MARKERS = {"", "-", "--", "—", "无", "空", "休", "休息", "空缺", "na", "n/a", "none", "null"}

# 员工分隔符：空格、逗号（含全角已折半）、顿号、分号、斜杠、竖线、换行
_TOKEN_SPLIT_RE = re.compile(r"[\s,、;/|]+")

# 工号形状：可选字母前缀 + 数字。前缀是否合法交给别名表判断——配置里工号可能是 S1、A007
_EMP_TOKEN_RE = re.compile(r"^([a-z]*)0*([0-9]{1,4})$")


def _split_tokens(cell: object) -> List[str]:
    return [t for t in _TOKEN_SPLIT_RE.split(_half("" if cell is None else str(cell)).strip()) if t]


# ---------- 配置方言（日期 / 班次 / 工号别名表） ----------


class _Dialect:
    """当前配置下「表格里的字面值 → 配置 id」的全部映射。

    单独成类而不是散落成模块常量，是因为别名表要随配置变化；而解析过程要反复查它，
    所以按 ConfigIndex 缓存一份（见 `_dialect_of`），不在每行里重建。
    """

    def __init__(self, index: ConfigIndex) -> None:
        self.index = index
        self.expected_slots = index.slot_count()
        self.days = _day_aliases(index)
        self.shifts = _shift_aliases(index)
        self.employee_exact, self.employee_numeric, self.id_prefixes = _employee_aliases(index)
        self.pool_desc = _pool_desc(index)
        # 有姓名字段时提示「精确匹配」，没有时沿用旧文案：店长看到的这句话决定他下一步怎么改表
        self.named = any(e.name != e.id for e in index.config.employees)
        self.reason_not_id = (
            f"不是合法工号（员工档案为 {self.pool_desc}，只按工号或完整姓名精确匹配）"
            if self.named
            else f"不是合法工号（员工档案只有 {self.pool_desc}，且无姓名字段）"
        )

    # ---- 查询 ----

    def day(self, raw: object) -> Optional[str]:
        return self.days.get(_norm(raw))

    def shift(self, raw: object) -> Optional[str]:
        return self.shifts.get(_norm(raw))

    def where(self, day: str, shift: str) -> str:
        return self.index.slot_label(day, shift)

    def resolve(self, token: str) -> Tuple[Optional[str], Optional[str]]:
        """把一个原始 token 归一为工号。返回 (工号, 失败原因)，两者必有一个为 None。

        容错到 `E1` / `e01` / `Ｅ０１` / `01` / `1` 是因为店长的表格是手打的；但容错止步于
        「形状能确定」的写法——超出档案或根本不是工号的一律带原因返回，交给上层进
        unresolved。这里刻意不做模糊匹配：一次错误的猜测会让整张体检报告的结论失真。
        """
        cleaned = _half(token).strip().strip("()[]{}（）:：,.")
        low = cleaned.lower()
        if low in _EMPTY_MARKERS:
            return None, None
        hit = self.employee_exact.get(_norm(cleaned))
        if hit:
            return hit, None
        m = _EMP_TOKEN_RE.match(low)
        if not m:
            return None, self.reason_not_id
        prefix, digits = m.group(1), m.group(2)
        # 前缀必须是档案里真实用过的（默认档案是 e）；「X21」不是「超出范围」，它根本不是工号
        if prefix and prefix not in self.id_prefixes:
            return None, self.reason_not_id
        eid = self.employee_numeric.get(int(digits))
        if eid is None:
            return None, f"工号 {cleaned} 超出员工档案范围 {self.pool_desc}"
        return eid, None


def _day_aliases(index: ConfigIndex) -> Dict[str, str]:
    """日期别名表。先登记 id/label 精确值，再补口语别名，保证精确值永远赢。"""
    out: Dict[str, str] = {}
    for d in index.day_ids:
        out.setdefault(_norm(d), d)
        out.setdefault(_norm(index.day_label(d)), d)
    for i, d in enumerate(index.day_ids):
        keys: List[str] = [f"周{d}", f"星期{d}", f"礼拜{d}", str(i + 1), f"day{i + 1}", f"d{i + 1}"]
        if len(index.day_ids) <= len(_EN_DAYS):
            # 英文星期按序号对应：只在周期不超过 7 天时成立，14 天周期里 mon 是哪一天并不确定
            keys += list(_EN_DAYS[i])
        if d == "日" or index.day_label(d).endswith("日"):
            keys += ["天", "周天", "星期天", "礼拜天"]
        for k in keys:
            out.setdefault(_norm(k), d)
    return out


def _shift_aliases(index: ConfigIndex) -> Dict[str, str]:
    out: Dict[str, str] = {}
    for s in index.shift_ids:
        out.setdefault(_norm(s), s)
        out.setdefault(_norm(index.shift_name(s)), s)
    for s in index.shift_ids:
        for key in (s, index.shift_name(s)):
            for alias in _SHIFT_SYNONYMS.get(key, ()):
                out.setdefault(_norm(alias), s)
    return out


def _employee_aliases(index: ConfigIndex) -> Tuple[Dict[str, str], Dict[int, str], set]:
    """工号别名表。

    停用员工也参与归一：导入的是**历史**排班表，读不懂离职员工等于读不懂这张表；
    是否该继续给他排班由校验器判断，不是解析层的事。
    """
    exact: Dict[str, str] = {}
    numeric: Dict[int, str] = {}
    ambiguous: set = set()
    prefixes: set = set()
    ids = [e.id for e in index.config.employees]
    for eid in ids:
        exact.setdefault(_norm(eid), eid)
        m = _EMP_TOKEN_RE.match(_norm(eid))
        if not m:
            continue
        prefixes.add(m.group(1))
        n = int(m.group(2))
        # 同一个数字对应多个工号（S1 与 A1）时放弃数字简写：宁可让店长写全，也不猜
        if n in numeric and numeric[n] != eid:
            ambiguous.add(n)
        numeric.setdefault(n, eid)
    for n in ambiguous:
        numeric.pop(n, None)
    # 姓名精确匹配：配置真的填了姓名时，「张三」是确定性映射而不是猜测；重名则一律不认
    seen_names: Dict[str, Optional[str]] = {}
    for e in index.config.employees:
        if e.name == e.id:
            continue
        key = _norm(e.name)
        seen_names[key] = None if key in seen_names else e.id
    for key, eid in seen_names.items():
        if eid and key not in exact:
            exact[key] = eid
    return exact, numeric, prefixes


def _pool_desc(index: ConfigIndex) -> str:
    """员工档案的可读范围描述，用于「超出范围」文案。默认配置下必须是 E01–E20。"""
    ids = [e.id for e in index.config.employees]
    if not ids:
        return "空（配置里没有员工）"
    parsed = [_EMP_TOKEN_RE.match(_norm(i)) for i in ids]
    if all(parsed) and len({m.group(1) for m in parsed if m}) == 1:
        nums = [int(m.group(2)) for m in parsed if m]
        if len(set(nums)) == len(nums) and max(nums) - min(nums) == len(nums) - 1:
            order = sorted(range(len(ids)), key=lambda k: nums[k])
            return f"{ids[order[0]]}–{ids[order[-1]]}"
    head = "、".join(ids[:4])
    return head if len(ids) <= 4 else f"{head} 等 {len(ids)} 个工号"


# 按索引缓存方言：ConfigIndex 本身已按配置指纹缓存，弱引用表让配置换掉后方言自然回收
_DIALECTS: "weakref.WeakKeyDictionary[ConfigIndex, _Dialect]" = weakref.WeakKeyDictionary()


def _dialect_of(config: ConfigLike = None) -> _Dialect:
    index = index_of(config)
    hit = _DIALECTS.get(index)
    if hit is None:
        hit = _Dialect(index)
        _DIALECTS[index] = hit
    return hit


# 默认维度的兼容出口：老调用方（含前端文案与测试）按 14 格理解「完整一周」
EXPECTED_SLOTS = default_index().slot_count()


def resolve_employee(token: str, config: ConfigLike = None) -> Tuple[Optional[str], Optional[str]]:
    """模块级归一入口，保持旧签名；语义见 `_Dialect.resolve`。"""
    return _dialect_of(config).resolve(token)


# ---------- 布局探测 ----------

_Row = Tuple[int, List[str]]   # (原始行号, 单元格) —— 行号只为 warning 可读


def _clean_rows(raw: List[List[object]]) -> List[_Row]:
    rows: List[_Row] = []
    for i, cells in enumerate(raw):
        vals = ["" if c is None else str(c).strip() for c in cells]
        if not any(vals):
            continue
        if vals[0].startswith("#"):      # 模板自带的注释行
            continue
        rows.append((i + 1, vals))
    return rows


def _detect_layout(rows: List[_Row], dialect: _Dialect) -> Tuple[Optional[int], str, dict]:
    """探测表头位置与布局。返回 (表头行号索引, layout, 列映射)。

    只扫前 6 行：真实文件常在最前面塞标题或说明，但表头不会掉到第 7 行才出现。
    """
    for idx, (_, cells) in enumerate(rows[:6]):
        norm = [_norm(c) for c in cells]
        # 首列已经是日期值（周一/mon/1）说明这是数据行，不是表头。少了这一步，
        # 无表头长表「周一,早班,E01 E06」会被误判成矩阵表，把「早班」当成人名读
        if norm and norm[0] in dialect.days:
            continue
        day_col = next((i for i, c in enumerate(norm) if c in _DAY_COLUMN_ALIASES), None)
        shift_col = next((i for i, c in enumerate(norm) if c in _SHIFT_COLUMN_ALIASES), None)
        emp_col = next((i for i, c in enumerate(norm) if c in _EMPLOYEE_COLUMN_ALIASES), None)
        shift_cols = {i: dialect.shifts[c] for i, c in enumerate(norm) if c in dialect.shifts}
        # 先判矩阵：矩阵表的班次名出现在列名上，比长表的三列特征更强
        if shift_cols and (day_col is not None or 0 not in shift_cols):
            return idx, "matrix", {"day": day_col if day_col is not None else 0, "shifts": shift_cols}
        if day_col is not None and shift_col is not None and emp_col is not None:
            return idx, "long", {"day": day_col, "shift": shift_col, "emp": emp_col}
    return None, "unknown", {}


def _detect_headerless(rows: List[_Row], dialect: _Dialect) -> Tuple[str, dict]:
    """没有表头时按首行形状推断，仍然是确定性判断：第二列是班次别名就是长表，否则是矩阵。"""
    if not rows:
        return "unknown", {}
    cells = rows[0][1]
    if _norm(cells[0]) not in dialect.days:
        return "unknown", {}
    if len(cells) >= 3 and _norm(cells[1]) in dialect.shifts:
        return "long", {"day": 0, "shift": 1, "emp": 2}
    if len(cells) >= 3:
        # 无表头矩阵：按配置班次顺序占据第 1..N 列，多出来的列没有班次可归，直接不看
        shifts = {i + 1: s for i, s in enumerate(dialect.index.shift_ids) if i + 1 < len(cells)}
        if shifts:
            return "matrix", {"day": 0, "shifts": shifts}
    return "unknown", {}


# ---------- 表格解析 ----------


class _Collector:
    """按 (day, shift) 累积工号，同时记账未匹配 token 与警告。"""

    def __init__(self, dialect: _Dialect) -> None:
        self.dialect = dialect
        self.grid: Dict[Tuple[str, str], List[str]] = {}
        self.unresolved: List[UnresolvedToken] = []
        self.warnings: List[str] = []
        self.assignments = 0
        self.resolved = 0

    def add_cell(self, day: str, shift: str, cell: object) -> None:
        where = self.dialect.where(day, shift)
        bucket = self.grid.setdefault((day, shift), [])
        for token in _split_tokens(cell):
            eid, reason = self.dialect.resolve(token)
            if reason is None and eid is None:
                continue                       # 「休」「-」这类占位写法，不是错误
            self.assignments += 1
            if eid is None:
                self.unresolved.append(
                    UnresolvedToken(raw=token, where=where, reason=reason or self.dialect.reason_not_id)
                )
                continue
            self.resolved += 1
            if eid in bucket:
                self.warnings.append(f"{where} 中 {eid} 重复出现，已去重")
                continue
            bucket.append(eid)

    def slots(self) -> List[Slot]:
        # 固定按配置的格子顺序输出（天优先），让前端网格渲染与 diff 都不必再排序
        return [
            Slot(day=d, shift=s, employee_ids=self.grid[(d, s)])
            for d, s in self.dialect.index.slots
            if (d, s) in self.grid
        ]


def parse_table_rows(raw_rows: List[List[object]], source: str, config: ConfigLike = None) -> ImportExtraction:
    """CSV / Excel 共用的解析主体：两者只有「怎么变成二维数组」这一步不同。"""
    dialect = _dialect_of(config)
    rows = _clean_rows(raw_rows)
    header_idx, layout, mapping = _detect_layout(rows, dialect)
    col = _Collector(dialect)

    if layout == "unknown":
        layout, mapping = _detect_headerless(rows, dialect)
        header_idx = None
        if layout != "unknown":
            col.warnings.append("未识别到表头，已按首行形状推断布局，建议使用「下载模板」的表头")

    if layout == "unknown":
        cols = " + ".join(dialect.index.shift_name(s) for s in dialect.index.shift_ids)
        return _failed_table(
            source,
            "没有识别到可用的表头：长表需要「日期 / 班次 / 员工」三列，"
            f"矩阵表需要「日期 + {cols}」列。可先下载 CSV 模板对照格式。",
        )

    body = rows if header_idx is None else rows[header_idx + 1:]
    seen: set = set()
    for lineno, cells in body:
        day_raw = cells[mapping["day"]] if mapping["day"] < len(cells) else ""
        day = dialect.day(day_raw)
        if day is None:
            if _norm(day_raw):
                col.warnings.append(f"第 {lineno} 行日期「{day_raw}」无法识别，已跳过")
            continue

        if layout == "long":
            shift_raw = cells[mapping["shift"]] if mapping["shift"] < len(cells) else ""
            shift = dialect.shift(shift_raw)
            if shift is None:
                col.warnings.append(f"第 {lineno} 行班次「{shift_raw}」无法识别，已跳过")
                continue
            if (day, shift) in seen:
                col.warnings.append(f"{dialect.where(day, shift)} 出现多行，已合并")
            seen.add((day, shift))
            # 员工列之后的所有列都当员工：有些表把 4 个人拆成 4 列
            col.add_cell(day, shift, " ".join(cells[mapping["emp"]:]))
        else:
            for ci, shift in mapping["shifts"].items():
                if (day, shift) in seen:
                    col.warnings.append(f"{dialect.where(day, shift)} 出现多行，已合并")
                seen.add((day, shift))
                col.add_cell(day, shift, cells[ci] if ci < len(cells) else "")

    return _finish(col, source=source, extractor="deterministic", layout=layout, model_used=None)


def _finish(
    col: _Collector, *, source: str, extractor: str, layout: str, model_used: Optional[str]
) -> ImportExtraction:
    dialect = col.dialect
    expected = dialect.expected_slots
    slots = col.slots()
    warnings = list(dict.fromkeys(col.warnings))
    if slots and col.resolved == 0:
        warnings.insert(
            0, f"识别到 {len(slots)} 个班次格子，但没有任何 token 能归一为 {dialect.pool_desc} 的合法工号"
        )
    ok = bool(slots) and col.resolved > 0
    if len(slots) != expected and ok:
        warnings.append(
            f"只解析出 {len(slots)} 个班次，完整周排班应有 {expected} 个，缺失班次会按 0 人参与校验"
        )

    if extractor == "vision_llm":
        # 视觉识别的可信度用「读到几格」×「工号匹配率」折算，直接摆给店长看
        match_rate = (col.resolved / col.assignments) if col.assignments else 1.0
        coverage = min(1.0, len(slots) / expected) if expected else 0.0
        confidence = round(coverage * match_rate, 2)
        # 图片来源恒为 true：视觉识别读错一个工号会让整份体检结论失真，
        # 而店长几乎不可能靠肉眼从校验报告里发现「基线本身读错了」
        requires_confirmation = True
    else:
        confidence = 1.0 if ok else 0.0
        requires_confirmation = bool(col.unresolved) or len(slots) != expected

    return ImportExtraction(
        ok=ok,
        source=source,  # type: ignore[arg-type]
        extractor=extractor,  # type: ignore[arg-type]
        layout=layout,  # type: ignore[arg-type]
        model_used=model_used,
        slots=slots,
        assignments=col.assignments,
        resolved=col.resolved,
        unresolved=col.unresolved,
        warnings=warnings,
        confidence=confidence,
        requires_confirmation=requires_confirmation,
    )


def _failed_table(source: str, message: str) -> ImportExtraction:
    return ImportExtraction(
        ok=False,
        source=source,  # type: ignore[arg-type]
        extractor="deterministic",
        layout="unknown",
        slots=[],
        warnings=[message],
        confidence=0.0,
        requires_confirmation=True,
    )


# ---------- CSV / TSV ----------

_DELIMITERS = [",", "\t", ";", "|"]


def _decode(data: bytes) -> Tuple[str, Optional[str]]:
    """按 utf-8-sig → gb18030 顺序解码：Windows 版 Excel 导出的 CSV 基本都是 GBK 系。"""
    for enc in ("utf-8-sig", "gb18030"):
        try:
            return data.decode(enc), None
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace"), "文件编码无法确定，已按 UTF-8 强制解码，部分字符可能失真"


def _pick_delimiter(text: str, dialect: _Dialect) -> Tuple[List[List[str]], str]:
    """用「解析后能否探测到布局」来选分隔符，而不是数字符出现次数。

    数次数会被单元格内的分隔符骗到：`E01;E02` 里的分号数量能压过真正的列分隔符逗号。
    """
    best: Optional[Tuple[tuple, List[List[str]], str]] = None
    for d in _DELIMITERS:
        rows = [list(r) for r in csv.reader(io.StringIO(text), delimiter=d)]
        _, layout, _ = _detect_layout(_clean_rows(rows), dialect)  # type: ignore[arg-type]
        score = (1 if layout != "unknown" else 0, max((len(r) for r in rows), default=0))
        if best is None or score > best[0]:
            best = (score, rows, d)
    assert best is not None
    return best[1], best[2]


def parse_csv(data: bytes, config: ConfigLike = None) -> ImportExtraction:
    text, warn = _decode(data)
    rows, _ = _pick_delimiter(text, _dialect_of(config))
    out = parse_table_rows(rows, "csv", config)  # type: ignore[arg-type]
    if warn:
        out.warnings.insert(0, warn)
    return out


# ---------- Excel ----------


def parse_excel(data: bytes, config: ConfigLike = None) -> ImportExtraction:
    """只用 openpyxl 读 .xlsx。

    .xls 是 BIFF 二进制老格式，openpyxl 明确不支持，而引入 xlrd 只为读一种正在被微软
    淘汰的格式并不划算——这里直接给出「另存为 .xlsx」的可执行提示，比抛异常有用。
    """
    try:
        import openpyxl
    except ImportError:      # pragma: no cover - 依赖已写进 requirements
        return _failed_table("excel", "服务端缺少 openpyxl 依赖，暂时无法解析 Excel，请改用 CSV 导入")

    try:
        wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    except Exception:
        return _failed_table(
            "excel",
            "无法读取该 Excel 文件（.xls 旧格式与加密文件都不支持）。"
            "请在 Excel 里「另存为 .xlsx」或导出 CSV 后重试。",
        )
    try:
        sheet_names = list(wb.sheetnames)
        ws = wb[sheet_names[0]]
        rows: List[List[object]] = []
        for row in ws.iter_rows(values_only=True):
            rows.append([_excel_cell(v) for v in row])
    finally:
        wb.close()
    out = parse_table_rows(rows, "excel", config)
    if len(sheet_names) > 1:
        out.warnings.append(f"该文件含 {len(sheet_names)} 个工作表，只解析了第一个「{sheet_names[0]}」")
    return out


def _excel_cell(v: object) -> str:
    # Excel 把「1」存成 1.0，直接 str() 会变成 "1.0" 而匹配不上日期别名「1」
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return "" if v is None else str(v)


# ---------- 图片 ----------


async def parse_image(
    data: bytes, ext: str, model: Optional[str] = None, config: ConfigLike = None
) -> ImportExtraction:
    """调视觉模型识别图片排班表。模型只负责「读出格子里写了什么」。"""
    dialect = _dialect_of(config)
    data_url = f"data:{_MIME_BY_EXT.get(ext, 'image/png')};base64,{base64.b64encode(data).decode()}"
    # 省略 config 时不传该关键字：视觉调用是被大量 mock 的边界，保持旧签名可调用
    extra = {} if config is None else {"config": config}
    try:
        rows, used = await llm.read_schedule_image(data_url, model, **extra)
    except ValueError as exc:
        return _failed_image(model, f"视觉模型输出无法解析（{exc}），建议重试或改用 CSV/Excel 导入")
    except Exception as exc:
        return _failed_image(
            model,
            f"调用视觉模型失败（{type(exc).__name__}），可能是超时或网络异常；"
            "可重试，或改用 CSV/Excel 导入以获得确定性结果",
        )

    col = _Collector(dialect)
    for r in rows:
        day = dialect.day(r.get("day"))
        shift = dialect.shift(r.get("shift"))
        if day is None or shift is None:
            col.warnings.append(f"识别结果中的「{r.get('day')} / {r.get('shift')}」不是合法日期与班次，已跳过")
            continue
        emps = r.get("employees")
        col.add_cell(day, shift, " ".join(str(e) for e in emps) if isinstance(emps, list) else emps)

    out = _finish(col, source="image", extractor="vision_llm", layout="image", model_used=used)
    out.warnings.insert(0, "图片内容由视觉模型识别，可能读错工号，请逐格核对后再应用为基线")
    return out


def _failed_image(model: Optional[str], message: str) -> ImportExtraction:
    return ImportExtraction(
        ok=False,
        source="image",
        extractor="vision_llm",
        layout="image",
        model_used=model,
        slots=[],
        warnings=[message],
        confidence=0.0,
        requires_confirmation=True,
    )


# ---------- 上传校验与模板 ----------


def extension_of(filename: str) -> str:
    name = (filename or "").lower()
    return name[name.rfind("."):] if "." in name else ""


def check_upload(filename: str, size: int) -> Optional[str]:
    """上传前置校验。返回 None 表示通过，否则是可直接展示给用户的中文原因。"""
    ext = extension_of(filename)
    if not ext:
        return "无法识别文件类型，请上传 .csv / .tsv / .xlsx / .png / .jpg 等格式"
    if ext not in ALLOWED_EXTS:
        return f"不支持的文件类型 {ext}。支持：{' '.join(sorted(ALLOWED_EXTS))}"
    if size <= 0:
        return "文件内容为空"
    if size > MAX_UPLOAD_BYTES:
        return f"文件 {size / 1024 / 1024:.2f} MB 超过 {MAX_UPLOAD_BYTES // 1024 // 1024} MB 上限，请压缩或改用 CSV 导入"
    return None


def _template_header(dialect: _Dialect) -> List[str]:
    index = dialect.index
    days = index.day_ids
    span = f"{index.day_label(days[0])}～{index.day_label(days[-1])}" if len(days) > 1 else \
        (index.day_label(days[0]) if days else "")
    sample = days[0] if days else ""
    shifts = " / ".join(index.shift_name(s) for s in index.shift_ids)
    name_hint = "系统只按工号或完整姓名精确匹配，写别名会进「未匹配」清单" if dialect.named \
        else "系统不会按姓名猜工号，填姓名会进「未匹配」清单"
    return [
        "# 智能排班助手 · 排班表导入模板（长表格式）",
        f"# 1) 日期填 {span}（也认「{sample}」「星期{sample}」「mon」）；班次填 {shifts}",
        "# 2) 员工列填工号，多个工号用空格、逗号、顿号或斜杠分隔",
        f"# 3) 只有 {dialect.pool_desc} 是合法工号；{name_hint}",
        f"# 4) 完整一个周期应有 {index.slot_count()} 行（{len(days)} 天 × {len(index.shift_ids)} 班）；"
        "以 # 开头的行会被忽略",
    ]


def template_csv(config: ConfigLike = None) -> str:
    """导入模板。示例内容取自求解器的合规解，保证模板本身零违规。

    模板里放一份「真的能过全部硬规则」的示例，是为了让店长第一次导入就看到 passed，
    而不是被一堆由模板自身违规造成的红叉劝退。
    """
    from . import solver
    from .models import ScheduleRequest

    dialect = _dialect_of(config)
    index = dialect.index
    schedule, _, _ = solver.solve(ScheduleRequest(), None, config=index)
    lines = _template_header(dialect)
    lines.append("日期,班次,员工")
    if schedule is None:      # 配置可能本身无解（自检会另行报错），兜底也要给出完整空表
        for d, s in index.slots:
            lines.append(f"{index.day_label(d)},{s},")
    else:
        for slot in schedule.slots:
            lines.append(f"{index.day_label(slot.day)},{slot.shift}," + " ".join(slot.employee_ids))
    return "\n".join(lines) + "\n"
